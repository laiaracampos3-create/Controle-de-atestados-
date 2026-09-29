from datetime import date, datetime, timedelta
import io
import json
import sqlite3
from google import genai
from google.genai import types
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
import pandas as pd
from PIL import Image
from pypdf import PdfReader
import streamlit as st

# --- Configurações da Página ---
st.set_page_config(
    page_title="Sistema de Gestão de Atestados | DP",
    page_icon="📋",
    layout="wide",
)

# Inicialização do cliente de IA (utiliza a variável de ambiente GEMINI_API_KEY do Render)
client = genai.Client()


# --- Gerenciamento do Banco de Dados ---
def init_db():
    conn = sqlite3.connect("atestados.db")
    cursor = conn.cursor()
    cursor.execute(
        """
        CREATE TABLE IF NOT EXISTS atestados (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            matricula TEXT,
            colaborador TEXT NOT NULL,
            setor TEXT,
            tipo_atestado TEXT,
            data_inicio DATE NOT NULL,
            dias INTEGER NOT NULL,
            data_fim DATE NOT NULL,
            cid TEXT,
            impacto_inss TEXT,
            observacoes TEXT
        )
    """
    )
    conn.commit()
    conn.close()


def salvar_atestado(
    matricula, colab, setor, tipo, data_ini, dias, cid, obs
):
    data_fim = data_ini + timedelta(days=int(dias) - 1)
    inss = "Sim" if int(dias) > 15 else "Não"

    conn = sqlite3.connect("atestados.db")
    cursor = conn.cursor()
    cursor.execute(
        """
        INSERT INTO atestados (
            matricula, colaborador, setor, tipo_atestado,
            data_inicio, dias, data_fim, cid, impacto_inss, observacoes
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """,
        (
            matricula,
            colab,
            setor,
            tipo,
            data_ini,
            dias,
            data_fim,
            cid,
            inss,
            obs,
        ),
    )
    conn.commit()
    conn.close()


def carregar_dados():
    conn = sqlite3.connect("atestados.db")
    df = pd.read_sql_query(
        "SELECT * FROM atestados ORDER BY data_inicio DESC", conn
    )
    conn.close()
    if not df.empty:
        df["data_inicio"] = pd.to_datetime(df["data_inicio"]).dt.date
        df["data_fim"] = pd.to_datetime(df["data_fim"]).dt.date
    return df


init_db()


# --- Extração com Inteligência Artificial Multimodal ---
def extrair_dados_arquivo(uploaded_file):
    nome_arq = uploaded_file.name.lower()
    prompt = """
    Você é um especialista em Departamento Pessoal e Medicina do Trabalho.
    Analise o documento do atestado médico (digital, escaneado ou foto manuscrita) e extraia rigorosamente em formato JSON:
    {
        "nome": "Nome completo do colaborador/paciente (ou vazio se ilegível)",
        "dias": número inteiro com a quantidade total de dias de afastamento (ex: repouso de 24h = 1, 48h = 2; declaração sem repouso = 0),
        "data_inicio": "AAAA-MM-DD (data de emissão ou início do repouso)",
        "cid": "Código do CID ou vazio"
    }
    Retorne apenas o JSON puro sem marcações extras.
    """

    conteudos = []

    # Se for imagem
    if nome_arq.endswith((".png", ".jpg", ".jpeg")):
        imagem = Image.open(uploaded_file)
        conteudos = [imagem, prompt]

    # Se for PDF
    elif nome_arq.endswith(".pdf"):
        reader = PdfReader(uploaded_file)
        texto = "".join([p.extract_text() or "" for p in reader.pages])

        if len(texto.strip()) > 30:
            conteudos = [prompt, f"Conteúdo do Atestado:\n{texto}"]
        else:
            uploaded_file.seek(0)
            bytes_pdf = uploaded_file.read()
            conteudos = [
                types.Part.from_bytes(
                    data=bytes_pdf, mime_type="application/pdf"
                ),
                prompt,
            ]

    try:
        resp = client.models.generate_content(
            model="gemini-2.5-flash",
            contents=conteudos,
            config=types.GenerateContentConfig(
                response_mime_type="application/json"
            ),
        )
        return json.loads(resp.text)
    except Exception as e:
        st.error(f"Erro na extração dos dados: {e}")
        return {}


# --- Geração de Planilha Profissional em Excel (OpenPyXL) ---
def gerar_planilha_fechamento_16a15(df_periodo, dt_inicio_corte, dt_fim_corte):
    output = io.BytesIO()

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        # Aba 1: Resumo Consolidado para Fechamento de Folha
        df_resumo = (
            df_periodo.groupby(
                ["matricula", "colaborador", "setor"], as_index=False
            )["dias_na_competencia"]
            .sum()
            .rename(
                columns={
                    "matricula": "Matrícula",
                    "colaborador": "Colaborador",
                    "setor": "Setor / Lotação",
                    "dias_na_competencia": "Total Dias a Abonar",
                }
            )
        )
        df_resumo.to_excel(
            writer, sheet_name="Resumo Folha de Ponto", index=False
        )

        # Aba 2: Detalhamento dos Registros
        df_detalhe = df_periodo[
            [
                "matricula",
                "colaborador",
                "setor",
                "tipo_atestado",
                "data_inicio",
                "data_fim",
                "dias",
                "dias_na_competencia",
                "cid",
                "impacto_inss",
                "observacoes",
            ]
        ].copy()

        df_detalhe.columns = [
            "Matrícula",
            "Colaborador",
            "Setor",
            "Tipo",
            "Início Atestado",
            "Fim Atestado",
            "Dias Totais",
            "Dias no Corte (16-15)",
            "CID",
            "Alerta INSS",
            "Observações",
        ]
        df_detalhe.to_excel(
            writer, sheet_name="Detalhamento Atestados", index=False
        )

        # Estilização visual executiva
        workbook = writer.book
        cor_cabecalho = PatternFill(
            start_color="1F4E78", end_color="1F4E78", fill_type="solid"
        )
        fonte_cabecalho = Font(
            name="Calibri", size=11, bold=True, color="FFFFFF"
        )
        fonte_dados = Font(name="Calibri", size=11)
        borda_fina = Border(
            left=Side(style="thin", color="E0E0E0"),
            right=Side(style="thin", color="E0E0E0"),
            top=Side(style="thin", color="E0E0E0"),
            bottom=Side(style="thin", color="E0E0E0"),
        )

        for sheetname in workbook.sheetnames:
            ws = workbook[sheetname]
            for col in range(1, ws.max_column + 1):
                cell = ws.cell(row=1, column=col)
                cell.fill = cor_cabecalho
                cell.font = fonte_cabecalho
                cell.alignment = Alignment(
                    horizontal="center", vertical="center"
                )

            for row in range(2, ws.max_row + 1):
                for col in range(1, ws.max_column + 1):
                    cell = ws.cell(row=row, column=col)
                    cell.font = fonte_dados
                    cell.border = borda_fina
                    if col in [1, 5, 6, 7, 8]:
                        cell.alignment = Alignment(
                            horizontal="center", vertical="center"
                        )

            for col in ws.columns:
                max_len = max(len(str(cell.value or "")) for cell in col)
                col_letter = col[0].column_letter
                ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    return output.getvalue()


# --- Interface Principal do Streamlit ---
st.title("📋 Painel de Controle de Atestados & Fechamento de Folha")
st.caption(
    "Extração automatizada de documentos médicos e apuração de competência 16 a 15."
)

tab_upload, tab_fechamento, tab_historico = st.tabs(
    [
        "📥 Novo Lançamento (Leitura IA)",
        "📊 Fechamento de Competência (16 a 15)",
        "📂 Base Histórica Geral",
    ]
)

# ----------------- ABA 1: LANÇAMENTO -----------------
with tab_upload:
    col_arq, col_form = st.columns([1, 1], gap="large")

    with col_arq:
        st.subheader("1. Envio do Documento")
        uploaded_file = st.file_uploader(
            "Arraste a foto ou PDF do atestado",
            type=["pdf", "png", "jpg", "jpeg"],
        )

        dados_lidos = {}
        if uploaded_file is not None:
            if uploaded_file.name.lower().endswith((".png", ".jpg", ".jpeg")):
                st.image(
                    uploaded_file,
                    caption="Prévia do documento enviado",
                    use_container_width=True,
                )

            # Controle do cache do arquivo em sessão
            file_key = f"proc_{uploaded_file.name}_{uploaded_file.size}"
            if st.session_state.get("last_uploaded") != file_key:
                with st.spinner("Realizando leitura inteligente do documento..."):
                    st.session_state.dados_lidos = extrair_dados_arquivo(
                        uploaded_file
                    )
                    st.session_state.last_uploaded = file_key

            dados_lidos = st.session_state.get("dados_lidos", {})
        else:
            st.session_state.pop("dados_lidos", None)
            st.session_state.pop("last_uploaded", None)

    with col_form:
        st.subheader("2. Validação e Lançamento")
        with st.form("form_registro", clear_on_submit=False):
            matricula = st.text_input("Matrícula")
            colaborador = st.text_input(
                "Nome Completo do Colaborador",
                value=dados_lidos.get("nome", ""),
            )
            setor = st.text_input("Setor / Lotação")

            c_tipo, c_cid = st.columns(2)
            tipo_atestado = c_tipo.selectbox(
                "Tipo de Documento",
                [
                    "Atestado Médico (Dias)",
                    "Declaração de Comparecimento (Horas)",
                    "Acidente de Trabalho",
                    "Licença Maternidade",
                    "Acompanhamento",
                ],
            )
            cid = c_cid.text_input("CID", value=dados_lidos.get("cid", ""))

            # Validação de data inicial
            data_ini_default = date.today()
            if dados_lidos.get("data_inicio"):
                try:
                    data_ini_default = datetime.strptime(
                        dados_lidos["data_inicio"], "%Y-%m-%d"
                    ).date()
                except Exception:
                    pass

            c_dt, c_dias = st.columns(2)
            data_inicio = c_dt.date_input(
                "Data de Início", value=data_ini_default
            )
            dias_val = int(dados_lidos.get("dias") or 1)
            dias = c_dias.number_input(
                "Dias de Afastamento",
                min_value=1,
                step=1,
                value=max(dias_val, 1),
            )

            observacoes = st.text_area(
                "Observações adicionais",
                placeholder="Ex: encaminhado para perícia médica, etc.",
            )

            gravar = st.form_submit_button(
                "Salvar Atestado no Banco", use_container_width=True
            )
            if gravar:
                if colaborador.strip():
                    salvar_atestado(
                        matricula,
                        colaborador,
                        setor,
                        tipo_atestado,
                        data_inicio,
                        dias,
                        cid,
                        observacoes,
                    )
                    st.success(
                        f"Atestado de {colaborador} registrado com sucesso!"
                    )
                else:
                    st.error("O campo 'Nome Completo' é obrigatório.")

# ----------------- ABA 2: FECHAMENTO 16 A 15 -----------------
with tab_fechamento:
    st.subheader("Apuração de Folha de Pagamento")

    hoje = date.today()
    c_m, c_a = st.columns(2)
    meses_pt = [
        "Janeiro",
        "Fevereiro",
        "Março",
        "Abril",
        "Maio",
        "Junho",
        "Julho",
        "Agosto",
        "Setembro",
        "Outubro",
        "Novembro",
        "Dezembro",
    ]
    mes_selecionado = c_m.selectbox(
        "Mês de Referência da Folha", meses_pt, index=hoje.month - 1
    )
    ano_selecionado = c_a.number_input(
        "Ano de Referência", min_value=2024, max_value=2035, value=hoje.year
    )

    mes_idx = meses_pt.index(mes_selecionado) + 1

    # Regra estrita de apuração: Dia 16 do mês anterior até Dia 15 do mês da folha
    if mes_idx == 1:
        dt_corte_inicio = date(ano_selecionado - 1, 12, 16)
    else:
        dt_corte_inicio = date(ano_selecionado, mes_idx - 1, 16)

    dt_corte_fim = date(ano_selecionado, mes_idx, 15)

    st.info(
        f"📅 **Período de Apuração da Folha:** {dt_corte_inicio.strftime('%d/%m/%Y')} a {dt_corte_fim.strftime('%d/%m/%Y')}"
    )

    df_base = carregar_dados()

    if not df_base.empty:
        # Atestados que tocam o período
        mask = (df_base["data_inicio"] <= dt_corte_fim) & (
            df_base["data_fim"] >= dt_corte_inicio
        )
        df_periodo = df_base[mask].copy()

        if not df_periodo.empty:

            def ratear_dias(row):
                ini = max(row["data_inicio"], dt_corte_inicio)
                fim = min(row["data_fim"], dt_corte_fim)
                return (fim - ini).days + 1

            df_periodo["dias_na_competencia"] = df_periodo.apply(
                ratear_dias, axis=1
            )

            # Métricas
            m1, m2, m3 = st.columns(3)
            m1.metric("Atestados no Período", len(df_periodo))
            m2.metric(
                "Total Dias a Abonar",
                int(df_periodo["dias_na_competencia"].sum()),
            )
            m3.metric(
                "Afastamentos > 15 dias (INSS)",
                len(df_periodo[df_periodo["impacto_inss"] == "Sim"]),
            )

            st.markdown("#### Resumo por Colaborador")
            resumo_vis = (
                df_periodo.groupby(
                    ["matricula", "colaborador", "setor"], as_index=False
                )["dias_na_competencia"]
                .sum()
                .rename(
                    columns={
                        "matricula": "Matrícula",
                        "colaborador": "Colaborador",
                        "setor": "Setor",
                        "dias_na_competencia": "Total Dias Abonados no Corte",
                    }
                )
            )
            st.dataframe(resumo_vis, use_container_width=True)

            # Geração do arquivo para download
            planilha_bytes = gerar_planilha_fechamento_16a15(
                df_periodo, dt_corte_inicio, dt_corte_fim
            )

            st.download_button(
                label=f"📥 Baixar Relatório de Fechamento ({mes_selecionado}/{ano_selecionado}) em Excel",
                data=planilha_bytes,
                file_name=f"Fechamento_Atestados_{mes_selecionado}_{ano_selecionado}_16a15.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
            )
        else:
            st.warning(
                "Nenhum atestado encontrado dentro deste período de apuração."
            )
    else:
        st.info("Nenhum dado registrado até o momento.")

# ----------------- ABA 3: HISTÓRICO GERAL -----------------
with tab_historico:
    st.subheader("Base Geral de Atestados Cadastrados")
    df_base = carregar_dados()
    if not df_base.empty:
        st.dataframe(df_base, use_container_width=True)
    else:
        st.write("Sem registros no momento.")
